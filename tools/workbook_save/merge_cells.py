"""Transplant only artifact-authored cells; preserve source XLSX native parts."""
from pathlib import Path
import zipfile,json,posixpath,copy,hashlib,sys,os
from lxml import etree as E
from openpyxl.utils.cell import range_boundaries,get_column_letter
W=Path(os.environ.get('SRDP_TASK_DIR') or Path(__file__).resolve().parent).resolve();N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda s:'{'+N+'}'+s
def package(p):
 with zipfile.ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}
def sheets(pkg):
 ids={r.get('Id'):posixpath.normpath(posixpath.join('xl',r.get('Target'))) if not r.get('Target').startswith('/') else r.get('Target').lstrip('/') for r in E.fromstring(pkg['xl/_rels/workbook.xml.rels'])}
 return {s.get('name'):ids[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')] for s in E.fromstring(pkg['xl/workbook.xml']).find(q('sheets'))}
for a in (list(map(int,sys.argv[1:])) or [1,2,5]):
 p=json.loads((W/f'A{a}_修改计划.json').read_text(encoding='utf-8'));target=Path(p['target'])
 assert hashlib.sha256(target.read_bytes()).hexdigest()==p['input_sha256'],'Concurrent source change'
 orig=package(target);out=dict(orig);patch=package(W/f'A{a}_authored_cells.xlsx');os=sheets(orig);ps=sheets(patch)
 strings=[''.join(t.itertext()) for t in E.fromstring(patch['xl/sharedStrings.xml'])] if 'xl/sharedStrings.xml' in patch else []
 trees={s:E.fromstring(orig[f]) for s,f in os.items()};pc={s:{c.get('r'):c for c in E.fromstring(patch[f]).iter(q('c'))} for s,f in ps.items()}
 for c in p['changes']:
  name=c['sheet'];root=trees[name];data=root.find(q('sheetData'));rows={int(r.get('r')):r for r in data};allcells={v.get('r'):v for v in root.iter(q('c'))}
  c0,r0,c1,r1=range_boundaries(c['range'])
  for rr in range(r0,r1+1):
   newrow=rr not in rows
   if newrow:
    row=E.SubElement(data,q('row'),r=str(rr));rows[rr]=row
    if c.get('style_template_row') is not None:
     template=rows.get(c['style_template_row'])
     if template is not None:
      for k,v in template.attrib.items():
       if k!='r':row.set(k,v)
   row=rows[rr];cells={v.get('r'):v for v in row}
   for cc in range(c0,c1+1):
    addr=f'{get_column_letter(cc)}{rr}';cell=cells.get(addr)
    if cell is None:
     cell=E.SubElement(row,q('c'),r=addr)
     tmpl=allcells.get(f'{get_column_letter(cc)}{c.get("style_template_row") or 6}')
     if tmpl is not None and tmpl.get('s'):cell.set('s',tmpl.get('s'))
    for el in list(cell):cell.remove(el)
    cell.attrib.pop('t',None);src=pc[name].get(addr)
    if src is not None:
     t=src.get('t')
     if t=='s':
      cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,q('is')),q('t')).text=strings[int(src.find(q('v')).text)]
     else:
      if t:cell.set('t',t)
      for v in src:
       if v.tag in [q('v'),q('f'),q('is')]:cell.append(copy.deepcopy(v))
   row[:]=sorted(row,key=lambda x:range_boundaries(x.get('r'))[0])
   if newrow:row.set('ht','36');row.set('customHeight','1')
  data[:]=sorted(data,key=lambda x:int(x.get('r')))
  dim=root.find(q('dimension'))
  if dim is not None:
   aa,bb,cc,dd=range_boundaries(dim.get('ref'));dim.set('ref',f'A1:{get_column_letter(max(cc,c1))}{max(dd,r1)}')
 for spec in p.get('styles',[]):
  root=trees[spec['sheet']];cells={v.get('r'):v for v in root.iter(q('c'))};style=next(v for v in trees[spec.get('from_sheet',spec['sheet'])].iter(q('c')) if v.get('r')==spec['from']).get('s');c0,r0,c1,r1=range_boundaries(spec['range'])
  for rr in range(r0,r1+1):
   for cc in range(c0,c1+1):
    dst=cells.get(f'{get_column_letter(cc)}{rr}')
    if dst is not None:
     if style is None:dst.attrib.pop('s',None)
     else:dst.set('s',style)
 for spec in p.get('column_widths',[]):
  root=trees[spec['sheet']];cols=root.find(q('cols'));template=next(v for v in cols if int(v.get('min'))<=spec['from']<=int(v.get('max')))
  for cc in range(spec['min'],spec['max']+1):
   present=next((v for v in cols if int(v.get('min'))<=cc<=int(v.get('max'))),None)
   if present is not None and int(present.get('min'))!=int(present.get('max')):raise ValueError('cannot overwrite grouped source columns')
   if present is None:present=E.SubElement(cols,q('col'),min=str(cc),max=str(cc))
   present.set('width',template.get('width'));present.set('customWidth','1')
  cols[:]=sorted(cols,key=lambda v:int(v.get('min')))
 for spec in p.get('row_heights',[]):
  data=trees[spec['sheet']].find(q('sheetData'));row=next(v for v in data if int(v.get('r'))==spec['row']);row.set('ht',str(spec['height']));row.set('customHeight','1')
 touched={spec['sheet'] for kind in ('changes','styles','row_heights','column_widths') for spec in p.get(kind,[])}
 for s in touched:out[os[s]]=E.tostring(trees[s],encoding='utf-8')
 changed=[n for n in orig if orig[n]!=out[n]];assert set(changed)<=set(os[s] for s in touched)
 with zipfile.ZipFile(W/f'A{a}_保存前.xlsx','w',zipfile.ZIP_DEFLATED) as z:
  for n,b in out.items():z.writestr(n,b)
 (W/f'A{a}_包检查.json').write_text(json.dumps({'changed_parts':changed,'all_other_parts_identical':True,'images_preserved':True},indent=2),encoding='utf-8')
 print(a,'draft ready, native objects preserved')
