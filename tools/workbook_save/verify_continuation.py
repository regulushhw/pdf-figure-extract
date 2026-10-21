"""Check plans and workbook candidates; --prepare-source alone records a source snapshot."""
from pathlib import Path
import argparse,copy,hashlib,io,json,os,sys,uuid,zipfile
import openpyxl
from openpyxl.utils.cell import range_boundaries
W=Path(os.environ.get('SRDP_TASK_DIR') or Path(__file__).resolve().parent).resolve()

def plan_errors(plan,book):
 errors=[]
 for change in plan['changes']:
  sheet=change['sheet']
  if sheet not in book.sheetnames:
   errors.append(['missing_sheet',sheet]);continue
  c0,r0,c1,r1=range_boundaries(change['range']);values=change['values']
  if len(values)!=r1-r0+1 or any(len(row)!=c1-c0+1 for row in values):
   errors.append(['range_shape',sheet,change['range']]);continue
  for rr,row in enumerate(values,r0):
   for cc,value in enumerate(row,c0):
    if value is None or value=='':continue
    for merged in book[sheet].merged_cells.ranges:
     if merged.min_row<=rr<=merged.max_row and merged.min_col<=cc<=merged.max_col and (rr,cc)!=(merged.min_row,merged.min_col):
      errors.append(['merged_non_anchor',sheet,book[sheet].cell(rr,cc).coordinate,str(merged)])
 formats={}
 for spec in plan.get('styles',[]):
  sheet=spec['sheet'];source_sheet=spec.get('from_sheet',sheet)
  if sheet not in book.sheetnames or source_sheet not in book.sheetnames:
   errors.append(['style_sheet_missing',sheet,source_sheet]);continue
  expected=book[source_sheet][spec['from']].number_format
  c0,r0,c1,r1=range_boundaries(spec['range'])
  for rr in range(r0,r1+1):
   for cc in range(c0,c1+1):
    key=(sheet,rr,cc)
    if key in formats and formats[key]!=expected:
     errors.append(['conflicting_number_format',sheet,openpyxl.utils.get_column_letter(cc)+str(rr),formats[key],expected])
    formats[key]=expected
 return errors

def content_errors(plan,book):
 errors=plan_errors(plan,book)
 if errors:return errors
 for change in plan['changes']:
  c0,r0,_,_=range_boundaries(change['range']);sheet=book[change['sheet']]
  for rr,row in enumerate(change['values'],r0):
   for cc,expected in enumerate(row,c0):
    actual=sheet.cell(rr,cc).value
    if actual!=expected and not(expected=='' and actual is None):
     errors.append(['value',sheet.title,sheet.cell(rr,cc).coordinate,expected,actual])
 for spec in plan.get('styles',[]):
  sheet=book[spec['sheet']];expected=book[spec.get('from_sheet',spec['sheet'])][spec['from']].number_format
  c0,r0,c1,r1=range_boundaries(spec['range'])
  for row in sheet.iter_rows(min_row=r0,max_row=r1,min_col=c0,max_col=c1):
   for cell in row:
    if cell.number_format!=expected:errors.append(['number_format',sheet.title,cell.coordinate,expected,cell.number_format])
 return errors

def media_hashes(path):
 with zipfile.ZipFile(path) as z:
  return sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('xl/media/'))

def image_anchors(sheet):
 """Compare image geometry in its worksheet, including offsets and extent."""
 def marker(value):
  return None if value is None else tuple(getattr(value,k) for k in ('col','colOff','row','rowOff'))
 def pair(value,keys):
  return None if value is None else tuple(getattr(value,k) for k in keys)
 return [(type(i.anchor).__name__,marker(getattr(i.anchor,'_from',None)),
          marker(getattr(i.anchor,'to',None)),pair(getattr(i.anchor,'pos',None),('x','y')),
          pair(getattr(i.anchor,'ext',None),('cx','cy')),getattr(i.anchor,'editAs',None)) for i in sheet._images]

def preservation_errors(plan,book,source):
 """Planned cells may change; source sheet order, other content and images may not."""
 errors=[];allowed={}
 for change in plan['changes']:
  c0,r0,c1,r1=range_boundaries(change['range'])
  allowed.setdefault(change['sheet'],set()).update((r,c) for r in range(r0,r1+1) for c in range(c0,c1+1))
 if book.sheetnames!=source.sheetnames:errors.append(['sheet_order',source.sheetnames,book.sheetnames])
 for name in set(source.sheetnames)&set(book.sheetnames):
  original=source[name];current=book[name]
  for key in set(original._cells)|set(current._cells):
   if key in allowed.get(name,set()):continue
   before=original._cells.get(key);after=current._cells.get(key)
   before_value=None if before is None else before.value;after_value=None if after is None else after.value
   before_type=None if before_value is None else before.data_type;after_type=None if after_value is None else after.data_type
   if (before_value,before_type)!=(after_value,after_type):
    errors.append(['unplanned_content',name,openpyxl.utils.get_column_letter(key[1])+str(key[0]),before_value,after_value])
  if image_anchors(original)!=image_anchors(current):errors.append(['image_anchors',name])
 return errors

def prepare_backup(plan,raw):
 """Prepare the declared backup before expensive authoring; never overwrite it."""
 backup=Path(plan['backup']).resolve()
 if backup==Path(plan['target']).resolve():return [['backup_is_target',str(backup)]]
 try:
  backup.parent.mkdir(parents=True,exist_ok=True)
  if backup.exists():
   if hashlib.sha256(backup.read_bytes()).hexdigest()!=plan['input_sha256']:
    return [['backup_conflict',str(backup)]]
  else:
   with backup.open('xb') as stream:stream.write(raw)
 except OSError as exc:return [['backup_unavailable',str(backup),str(exc)]]
 return []


def prepare_source(paper,plan,plan_path):
 """Capture the preflight-passing original once; reuse only its matching snapshot."""
 plan_bytes=plan_path.read_bytes();raw=Path(plan['target']).read_bytes()
 if json.loads(plan_bytes)!=plan:return [['plan_changed','Retry from the current plan.']],None
 if hashlib.sha256(raw).hexdigest()!=plan['input_sha256']:
  return [['source_changed','Rebuild the plan from the current saved file.']],None
 source=openpyxl.load_workbook(io.BytesIO(raw),data_only=False)
 try:errors=plan_errors(plan,source)
 finally:source.close()
 if errors:return errors,None
 if plan_path.read_bytes()!=plan_bytes:return [['plan_changed','Retry from the current plan.']],None
 if plan.get('source_snapshot'):
  snapshot=Path(plan['source_snapshot'])
  if not snapshot.is_file() or hashlib.sha256(snapshot.read_bytes()).hexdigest()!=plan['input_sha256']:
   return [['source_snapshot_invalid','Existing snapshot is missing or changed; rebuild the plan from the correct original.']],None
  errors=prepare_backup(plan,raw)
  if errors:return errors,None
  return [],str(snapshot)
 errors=prepare_backup(plan,raw)
 if errors:return errors,None
 snapshot=W/f'A{paper}_原件快照_{uuid.uuid4().hex}.xlsx'
 with snapshot.open('xb') as stream:stream.write(raw)
 plan['source_snapshot']=str(snapshot)
 temporary=plan_path.with_name(plan_path.name+'.'+uuid.uuid4().hex+'.tmp')
 try:
  temporary.write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
  if plan_path.read_bytes()!=plan_bytes:raise ValueError('Plan changed while preparing source snapshot.')
  os.replace(temporary,plan_path)
 finally:
  if temporary.exists():temporary.unlink()
 return [],str(snapshot)

def verify(paper,draft=False,plan_only=False,require_current_source=False,candidate=None,prepare=False):
 plan_path=W/f'A{paper}_修改计划.json';plan=json.loads(plan_path.read_text(encoding='utf-8'))
 target=Path(candidate) if candidate else W/f'A{paper}_保存前.xlsx' if draft else Path(plan['target'])
 book=openpyxl.load_workbook(target,data_only=False)
 errors=plan_errors(plan,book) if plan_only else content_errors(plan,book)
 if require_current_source and hashlib.sha256(Path(plan['target']).read_bytes()).hexdigest()!=plan['input_sha256']:
  errors.append(['source_changed','Rebuild the plan from the current saved file.'])
 result={'stage':'plan' if plan_only else 'draft' if draft else 'candidate' if candidate else 'saved','path':str(target),'errors':errors,'all_sheets':[(s.title,s.max_row,s.max_column,len(s._images)) for s in book]}
 if prepare and not errors:
  prepared,snapshot=prepare_source(paper,plan,plan_path);errors.extend(prepared)
  if snapshot:result['source_snapshot']=snapshot
 if not plan_only:
  snapshot=Path(plan['source_snapshot']) if plan.get('source_snapshot') else None
  result['preservation_check']='not checked: no valid per-run source snapshot'
  if snapshot is None:errors.append(['source_snapshot_missing','Unchanged content was not checked; prepare a new plan before saving.'])
  elif not snapshot.is_file() or hashlib.sha256(snapshot.read_bytes()).hexdigest()!=plan['input_sha256']:
   errors.append(['source_snapshot_invalid','The snapshot must match input_sha256.'])
  else:
   source=openpyxl.load_workbook(snapshot,data_only=False)
   try:errors.extend(preservation_errors(plan,book,source))
   finally:source.close()
   result['preservation_check']='checked against per-run source snapshot'
   result['images_unchanged']=media_hashes(target)==media_hashes(snapshot)
   if not result['images_unchanged']:errors.append(['images_changed'])
  with zipfile.ZipFile(target) as z:
   if z.testzip() is not None:errors.append(['invalid_zip'])
  result['checked_written_cells']=sum(len(row) for c in plan['changes'] for row in c['values'])
  result['cache_check']='pending final engine save' if draft else 'checked'
  if not draft:
   cached=openpyxl.load_workbook(target,data_only=True)
   for s in book:
    for cell in s._cells.values():
     val=cached[s.title][cell.coordinate]
     if val.data_type=='e' or(cell.data_type=='f' and val.value is None):errors.append(['formula_cache',s.title,cell.coordinate])
   cached.close()
 book.close()
 return result

def static_engine_check(paper,engine_path):
 """Compare every typed value after full-engine export; no formula is allowed.

 Engine blank-string normalization is recorded, not applied to the original
 package. A caller may retain the native draft only after this check passes.
 """
 draft_path=W/f'A{paper}_保存前.xlsx'
 draft=openpyxl.load_workbook(draft_path,data_only=False)
 engine=openpyxl.load_workbook(engine_path,data_only=False)
 errors=[];checked=normalized=0
 try:
  if draft.sheetnames!=engine.sheetnames:errors.append(['sheet_order',draft.sheetnames,engine.sheetnames])
  for name in set(draft.sheetnames)&set(engine.sheetnames):
   before,after=draft[name],engine[name]
   for key in set(before._cells)|set(after._cells):
    left=before._cells.get(key);right=after._cells.get(key)
    lv=None if left is None else left.value;rv=None if right is None else right.value
    lt=None if lv is None else left.data_type;rt=None if rv is None else right.data_type
    cell=openpyxl.utils.get_column_letter(key[1])+str(key[0]);checked+=1
    if lt=='f' or rt=='f':errors.append(['formula_not_supported',name,cell]);continue
    if rt=='e':errors.append(['engine_error',name,cell,rv]);continue
    if lv in (None,'') and rv in (None,''):
     normalized+=int((lv,lt)!=(rv,rt));continue
    if (lv,lt)!=(rv,rt):errors.append(['engine_value',name,cell,lv,rv,lt,rt])
 finally:
  draft.close();engine.close()
 return {'stage':'full-engine static-value check','draft':str(draft_path),'engine_export':str(engine_path),
         'draft_sha256':hashlib.sha256(draft_path.read_bytes()).hexdigest(),
         'engine_sha256':hashlib.sha256(Path(engine_path).read_bytes()).hexdigest(),
         'checked_cells':checked,'blank_normalizations':normalized,'formula_policy':'reject any formula',
         'errors':errors}

def main():
 parser=argparse.ArgumentParser();parser.add_argument('papers',nargs='*',type=int)
 for flag in ['draft','plan-only','require-current-source','prepare-source','engine-values']:parser.add_argument('--'+flag,action='store_true')
 parser.add_argument('--candidate',type=Path)
 args=parser.parse_args()
 if args.engine_values:
  if len(args.papers)!=1 or args.candidate is None or args.draft or args.plan_only:parser.error('--engine-values requires one paper and --candidate full-engine-export.')
  result=static_engine_check(args.papers[0],args.candidate)
  (W/f'续补引擎核对_A{args.papers[0]}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
  print(f'A{args.papers[0]}: {len(result["errors"])} engine value errors; {result["blank_normalizations"]} blank normalizations retained only in the engine export')
  for error in result['errors'][:8]:print(json.dumps(error,ensure_ascii=False))
  return 1 if result['errors'] else 0
 if args.prepare_source and not(args.plan_only and args.require_current_source):parser.error('--prepare-source requires --plan-only and --require-current-source.')
 if args.candidate and (len(args.papers)!=1 or args.draft or args.plan_only):parser.error('--candidate requires exactly one paper and cannot be combined with --draft or --plan-only.')
 results={}
 for paper in args.papers or [1]:
  try:results[str(paper)]=verify(paper,args.draft,args.plan_only,args.require_current_source,args.candidate,args.prepare_source)
  except Exception as exc:results[str(paper)]={'errors':[['exception',str(exc)]]}
 if not args.plan_only:
  name=f'续补候选检查_A{args.papers[0]}.json' if args.candidate else '续补草稿检查.json' if args.draft else '续补保存检查.json'
  (W/name).write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
 for paper,result in results.items():
  print(f'A{paper}: {len(result["errors"])} errors ({result.get("stage","failed")})')
  for error in result['errors'][:12]:print(json.dumps(error,ensure_ascii=False))
 return 1 if any(r['errors'] for r in results.values()) else 0

if __name__=='__main__':sys.exit(main())
