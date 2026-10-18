# -*- coding: utf-8 -*-
"""底稿生成器：A1 版式模板。

版式要点（与项目内已交付底稿一致）：
- 数据页五列固定在前：对象 | 情景或模型条件 | 指标 | 所属时段 | 出处，数值横排向右。
- 单位并入数字格式（0.00"%" 等），不单列。
- 行4 留白(高8)，表头行高32，数据行34-38；微软雅黑，标题16粗，正文11。
- 数值单元格关闭自动换行（坑库 #7）。

用法（库）：
    wb = Workbook(); ws = wb.active
    r = head_block(ws, "标题", ["说明1", "说明2"], ncols=8,
                   widths=[24,21,20,20,20,14,14,14],
                   headers=[...])
    r = data_row(ws, r, ["对象","情景","指标","时段","出处"], [1.0, 2.0, 3.0], '0.00"%"')
    wb.save("底稿.xlsx")
"""
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

YH = 'Microsoft YaHei'
F_TITLE = Font(name=YH, size=16, bold=True)
F_NOTE = Font(name=YH, size=11)
F_HEAD = Font(name=YH, size=11, bold=True)
F_BODY = Font(name=YH, size=11)
AL_C = Alignment(wrap_text=True, vertical='center', horizontal='center')
AL_L = Alignment(wrap_text=True, vertical='center')
AL_R = Alignment(wrap_text=False, vertical='center', horizontal='right', indent=1)  # 坑库 #7
AL_T = Alignment(wrap_text=True, vertical='top')
UNDER = Border(bottom=Side(style='thin', color='D9D9D9'))


def head_block(ws, title, notes, ncols, widths, headers, head_h=32):
    """标题 + 说明 + 表头。返回首条数据行号（表头下一行）。"""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws['A1'] = title
    ws['A1'].font = F_TITLE
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 33
    r = 2
    for t in notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1, value=t)
        c.font = F_NOTE
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[r].height = 30 if len(t) < 60 else 36
        r += 1
    ws.row_dimensions[4].height = 8
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=i, value=h)
        c.font = F_HEAD
        c.alignment = AL_C
        c.border = UNDER
    ws.row_dimensions[5].height = head_h
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return 6


def data_row(ws, row, idents, vals, fmt, rh=38):
    """一行数据：idents 从第1列起，vals 数值横排（单位并入 fmt）。"""
    for i, v in enumerate(idents, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = F_BODY
        c.alignment = AL_L
    start = len(idents) + 1
    for j, v in enumerate(vals):
        c = ws.cell(row=row, column=start + j, value=v)
        c.font = F_BODY
        c.alignment = AL_R
        c.number_format = fmt
    ws.row_dimensions[row].height = rh
    return row + 1


def finalize_sheet(ws, first_data_row=6):
    """写完数据再设置筛选；不冻结窗格。"""
    ws.auto_filter.ref = f'A{first_data_row-1}:{get_column_letter(ws.max_column)}{ws.max_row}'
    ws.freeze_panes = None
