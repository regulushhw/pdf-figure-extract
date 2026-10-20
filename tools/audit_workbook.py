"""按原文预先建立的审查契约检查底稿；不证明契约本身正确。
python tools/audit_workbook.py workbook.xlsx audit.json
audit.json: {"rules":[{"kind":"cell","sheet":"数据","cell":"C6",
"expected":-0.21,"source":"PDF p5 Table2"}, ...]}
支持cell、count、sequence、source_records、inventory、extraction_records、sampling；失败退出1，空规则拒绝通过。
source_records的json相对契约文件，内容为[{sheet,cell,value,method/status},...]。
"""
import json,sys,math,numbers,hashlib
from pathlib import Path
import openpyxl
from verify import formula_cache_report

def inventory_check(ws, rule):
    """Compare source-declared series with rows, including completely absent series."""
    from openpyxl.utils.cell import column_index_from_string
    keys=rule['key_columns'];values=rule['value_columns']
    if not keys or not values or len(set(keys+values))!=len(keys+values):
        raise ValueError('身份列与数值列须非空、唯一且互不重叠')
    for col in keys+values:column_index_from_string(col)
    first,last=rule['row_range']
    if not isinstance(first,int) or not isinstance(last,int) or not 1<=first<=last:
        raise ValueError('无效行范围')
    expected={}
    for item in rule['expected']:
        key=tuple(item['key']);count=item['count']
        if len(key)!=len(keys) or key in expected:raise ValueError('预期身份键重复或长度错误')
        if type(count) is not int or not 0<count<=len(values):raise ValueError('预期点数无效')
        expected[key]=count
    if not expected:raise ValueError('原图系列清单为空')
    actual={};errors=[]
    for row in range(first,last+1):
        key=tuple(ws[f'{col}{row}'].value for col in keys)
        vals=[ws[f'{col}{row}'].value for col in values]
        finite=lambda v:isinstance(v,numbers.Real) and not isinstance(v,bool) and math.isfinite(float(v))
        if key not in expected and not any(finite(v) for v in vals):continue
        if key in actual:errors.append(f'重复系列：{key!r}（行{actual[key][0]}/{row}）')
        actual[key]=(row,sum(finite(v) for v in vals))
        if any(isinstance(v,numbers.Real) and not isinstance(v,bool) and not finite(v) for v in vals):
            errors.append(f'非有限值：行{row}')
    for key,count in expected.items():
        if key not in actual:errors.append(f'漏掉整条系列：{key!r}')
        elif actual[key][1]!=count:errors.append(f'系列点数{actual[key][1]}，预期{count}：{key!r}')
    for key in actual.keys()-expected.keys():errors.append(f'清单外系列：{key!r}')
    if errors:raise ValueError('; '.join(errors))
    return len(expected)

def same(a,b,tol=1e-8):
    if isinstance(a,numbers.Real) and isinstance(b,numbers.Real):
        return math.isfinite(float(a)) and math.isfinite(float(b)) and abs(a-b)<=tol
    return type(a) is type(b) and a==b

def extraction_records_check(wb, rule, base):
    """Check saved values against retained raw readings, not regenerated expectations."""
    path=base/rule['json']
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    if digest!=rule['sha256']:raise ValueError('原始取点记录摘要不符')
    rows=json.loads(path.read_text(encoding='utf-8-sig'))
    if not rows:raise ValueError('原始取点记录为空')
    seen=set()
    for r in rows:
        key=(r['sheet'],r['cell'])
        if key in seen:raise ValueError(f'重复取点记录：{key}')
        seen.add(key)
        raw=r['raw_value'];value=wb[key[0]][key[1]].value
        method=r['method']
        if not r.get('source'):raise ValueError(f'取点缺出处：{key}')
        finite=lambda v:type(v) in (int,float) and math.isfinite(v)
        if raw is not None and not finite(raw):raise ValueError(f'原始读数无效：{key}')
        if method in ('direct','neighbor'):
            if raw is None or not same(value,raw):
                raise ValueError(f'直接/邻列读数被替换：{key}')
        elif method=='missing':
            if raw is not None or value is not None:raise ValueError(f'缺值被填入：{key}')
        elif method=='interpolated':
            if raw is not None:raise ValueError(f'插值覆盖已有读数：{key}')
            if not finite(value) or not r.get('calculation'):
                raise ValueError(f'插值缺计算依据：{key}')
            if not same(value,r['calculation']['value']):raise ValueError(f'插值结果不符：{key}')
            if not r['calculation'].get('inputs'):raise ValueError(f'插值缺输入：{key}')
        elif method=='reviewed':
            review=r.get('review',{});pixel=review.get('pixel',[])
            if not review.get('source') or len(pixel)!=2 or not all(finite(v) for v in pixel):
                raise ValueError(f'重读缺原图与坐标：{key}')
            if not finite(value) or not same(value,review.get('value')):
                raise ValueError(f'重读结果不符：{key}')
        else:raise ValueError(f'未知取点方法：{method}')
    expected={(rule['sheet'],c.coordinate) for row in wb[rule['sheet']][rule['range']] for c in row}
    if seen!=expected:raise ValueError(f'取点覆盖不符：漏{len(expected-seen)}，越界{len(seen-expected)}')
    return len(seen)

def sampling_check(ws, rule):
    """Check agreed coordinates, including missing positions, against saved headers."""
    expected=rule['expected'];cells=[c for row in ws[rule['range']] for c in row]
    if not expected or any(v is None for v in expected):raise ValueError('约定取样坐标为空')
    if len(expected)!=len(set(expected)):raise ValueError('约定取样坐标重复')
    actual=[c.value for c in cells]
    if len(actual)!=len(expected) or not all(same(a,b) for a,b in zip(actual,expected)):
        raise ValueError('实际日期/概率取样与约定不符（漏点、错位或顺序错误）')
    return len(expected)

def audit(workbook,contract):
    contract=Path(contract);config=json.loads(contract.read_text(encoding='utf-8-sig'))
    rules=config.get('rules',[])
    if not rules:raise ValueError('无审查规则，不能通过')
    wb=openpyxl.load_workbook(workbook,data_only=True)
    formulas=openpyxl.load_workbook(workbook,data_only=False)
    cache_report=formula_cache_report(formulas,wb)
    failures=[{'rule':'formula_cache','error':message} for message in cache_report['failures']];checked=0
    for index,rule in enumerate(rules):
        try:
            if not rule.get('source'):raise ValueError('缺少原文/依据定位')
            kind=rule['kind']
            if kind=='inventory':
                checked+=inventory_check(wb[rule['sheet']],rule)
            elif kind=='extraction_records':
                checked+=extraction_records_check(wb,rule,contract.parent)
            elif kind=='sampling':
                checked+=sampling_check(wb[rule['sheet']],rule)
            elif kind=='source_records':
                rows=json.loads((contract.parent/rule['json']).read_text(encoding='utf-8-sig'))
                if not rows:raise ValueError('源记录为空')
                for r in rows:
                    got=wb[r['sheet']][r['cell']].value;want=r['value']
                    if not same(got,want):raise ValueError(f"源值与底稿不符：{r['sheet']}!{r['cell']}")
                    if want==0 and r.get('method') not in ('original_table','explicit_zero_height','baseline_visually_near_zero'):
                        raise ValueError(f"0缺少明确依据：{r['cell']}")
                    if r.get('status')=='missing' and got is not None:raise ValueError('缺值被填入数值')
                    checked+=1
            elif kind=='cell':
                got=wb[rule['sheet']][rule['cell']].value
                if not same(got,rule['expected']):raise ValueError(f"身份/单位/数值不符：{rule['cell']} got={got!r}")
                checked+=1
            elif kind in ('count','sequence'):
                cells=[c for row in wb[rule['sheet']][rule['range']] for c in row]
                if kind=='count':
                    n=sum(isinstance(c.value,numbers.Real) and not isinstance(c.value,bool) for c in cells)
                    if n!=rule['expected']:raise ValueError(f"数值覆盖{n}，预期{rule['expected']}")
                elif len(cells)!=len(rule['expected']) or not all(same(c.value,v) for c,v in zip(cells,rule['expected'])):
                    raise ValueError('坐标/日期标签不符')
                checked+=1
            else:raise ValueError(f'未知检查类型：{kind}')
        except Exception as exc:failures.append({'rule':index,'error':str(exc)})
    return {'checked':checked,'failures':failures,'formula_cells':cache_report['formula_cells'],
            'scope':'declared_contract_only; original_source_review_still_required'}

if __name__=='__main__':
    try:result=audit(sys.argv[1],sys.argv[2])
    except Exception as exc:result={'failures':[str(exc)]}
    print(json.dumps(result,ensure_ascii=False,indent=2))
    sys.exit(1 if result['failures'] else 0)

