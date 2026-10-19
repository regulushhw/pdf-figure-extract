# -*- coding: utf-8 -*-
"""核对写表一致性，不证明原图提取正确。
checks.json: {"ranges":[{"sheet":"数据","range":"F6:Q8"}],
             "checks":[{"sheet":"数据","cell":"F6","json":"figure.json",
                        "path":["series","A",0,1],"tol":0.001}]}
所有range中的单元格必须恰好检查一次；范围必须覆盖约定数值区，不能只挑易核项。
JSON路径相对于checks.json；列表形式path支持含小数点的键。
"""
import json,sys,math,numbers
from pathlib import Path
import openpyxl
from openpyxl.utils.cell import range_boundaries


def get_by_path(obj,path):
    for k in (path.split('.') if isinstance(path,str) else path):
        obj = obj[int(k)] if isinstance(obj,list) else obj[k]
    return obj


def formula_cache_report(formulas,values):
    """Read-only saved-file check, including formulas outside declared ranges.

    An empty cache is not proof of an intentionally blank formula result.
    Resolve such cases in the calculation engine before reporting success.
    This checks readability/errors, not formula arithmetic or freshness.
    """
    count=0;failures=[]
    for ws in formulas:
        for row in ws:
            for cell in row:
                if cell.data_type!='f':continue
                count+=1;cached=values[ws.title][cell.coordinate]
                label=f'{ws.title}!{cell.coordinate}'
                if cached.value is None:
                    failures.append(f'{label}: formula cache is empty; recalculate after the last edit')
                elif cached.data_type=='e':
                    failures.append(f'{label}: formula cache contains Excel error {cached.value}')
                elif isinstance(cached.value,numbers.Real) and not math.isfinite(float(cached.value)):
                    failures.append(f'{label}: formula cache is not finite')
    return {'formula_cells':count,'failures':failures}


def verify(xlsx,checks_path):
    base=Path(checks_path).resolve().parent
    config=json.loads(Path(checks_path).read_text(encoding='utf-8-sig'))
    if not isinstance(config,dict) or not config.get('checks') or not config.get('ranges'):
        raise ValueError('必须提供非空checks与完整数值区ranges；空检查不能通过')
    wb=openpyxl.load_workbook(xlsx,data_only=True)
    formulas=openpyxl.load_workbook(xlsx,data_only=False)
    cache_report=formula_cache_report(formulas,wb)
    expected=set()
    for r in config['ranges']:
        lo,top,hi,bottom=range_boundaries(r['range'])
        ws=wb[r['sheet']]
        for row in ws.iter_rows(min_row=top,max_row=bottom,min_col=lo,max_col=hi):
            for cell in row:expected.add((ws.title,cell.coordinate))
    if not expected:raise ValueError('检查范围为空')
    seen=set();failures=list(cache_report['failures'])
    for c in config['checks']:
        key=(c['sheet'],c['cell'].upper())
        if key in seen:raise ValueError(f'重复检查：{key}')
        seen.add(key)
        try:
            got=wb[key[0]][key[1]].value
            if formulas[key[0]][key[1]].data_type=='f' and got is None:
                raise ValueError('公式无可读缓存，须先重算')
            want=get_by_path(json.loads((base/c['json']).read_text(encoding='utf-8-sig')),c['path'])
            tol=float(c.get('tol',1e-6))
            if not math.isfinite(tol) or tol<0:raise ValueError('容差须非负且有限')
            if want is None:
                valid=got is None
            elif isinstance(want,numbers.Real) and not isinstance(want,bool):
                valid=(isinstance(got,numbers.Real) and not isinstance(got,bool)
                       and math.isfinite(float(got)) and math.isfinite(float(want))
                       and abs(got-want)<=tol)
            else:valid=type(got) is type(want) and got==want
            if not valid:raise ValueError(f'{key}: got={got!r}, want={want!r}')
        except Exception as exc:failures.append(str(exc))
    if seen!=expected:
        failures.append(f'覆盖不符：漏检{len(expected-seen)}，越界{len(seen-expected)}')
    return {'checked':len(seen),'expected':len(expected),'failures':failures,
            'formula_cells':cache_report['formula_cells'],
            'scope':'workbook_vs_json_only'}


def main():
    try:result=verify(sys.argv[1],sys.argv[2])
    except Exception as exc:
        print('FAIL:',exc);sys.exit(1)
    print(json.dumps(result,ensure_ascii=False,indent=2))
    sys.exit(1 if result['failures'] else 0)

if __name__=='__main__':main()

