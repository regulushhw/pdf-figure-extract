"""Synthetic save-entry regressions; optional configured artifact-engine integration."""
import copy
import hashlib
import io
import json
import os
from pathlib import Path
import runpy
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

TOOLS=Path(__file__).resolve().parents[1]/'tools/workbook_save'

class WorkbookSaveTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.book=openpyxl.Workbook();s=self.book.active;s.title='Data'
        s['A1']=1;s['B1']=2;s['B2']='merged';s.merge_cells('B2:C2')
        f=self.book.create_sheet('Formats');f['A1']=0;f['A1'].number_format='0.00'
        f['A2']=0;f['A2'].number_format='0.0'
        Image.new('RGB',(8,8),'blue').save(self.root/'image.png')
        s.add_image(SheetImage(io.BytesIO((self.root/'image.png').read_bytes())),'E5')
        self.source=self.root/'source.xlsx';self.book.save(self.source)
        self.plan={'paper':1,'target':str(self.source),'backup':str(self.root/'backup.xlsx'),
            'input_sha256':self.sha(self.source),
            'changes':[{'sheet':'Data','range':'A1','values':[[3]]}],
            'styles':[{'sheet':'Data','range':'A1','from_sheet':'Formats','from':'A1'}],
            'previews':[{'sheet':'Data','range':'A1:C3','label':'synthetic'}]}
        with patch.dict(os.environ,{'SRDP_TASK_DIR':str(self.root)}):
            self.mod=runpy.run_path(str(TOOLS/'verify_continuation.py'))
    @staticmethod
    def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
    def test_default_and_explicit_directory(self):
        self.assertEqual(self.mod['W'],self.root.resolve())
        with patch.dict(os.environ,{'SRDP_TASK_DIR':''}):
            self.assertEqual(runpy.run_path(str(TOOLS/'verify_continuation.py'))['W'],TOOLS.resolve())
    def test_conflicting_formats_rejected_before_authoring(self):
        p=copy.deepcopy(self.plan)
        p['styles'].append({'sheet':'Data','range':'A1','from_sheet':'Formats','from':'A2'})
        self.assertTrue(any(e[0]=='conflicting_number_format' for e in self.mod['plan_errors'](p,self.book)))
        p['styles'][-1]['from']='A1'
        self.assertEqual(self.mod['plan_errors'](p,self.book),[])
    def test_missing_style_sheet_and_merged_guard(self):
        p=copy.deepcopy(self.plan);p['styles'][0]['from_sheet']='missing'
        self.assertTrue(any(e[0]=='style_sheet_missing' for e in self.mod['plan_errors'](p,self.book)))
        p={'changes':[{'sheet':'Data','range':'C2','values':[['bad']]}]}
        self.assertTrue(any(e[0]=='merged_non_anchor' for e in self.mod['plan_errors'](p,self.book)))
    def test_unplanned_value_and_deleted_sheet_rejected(self):
        changed=openpyxl.load_workbook(self.source,data_only=False)
        try:
            changed['Data']['B1']=99
            self.assertTrue(any(e[0]=='unplanned_content' for e in self.mod['preservation_errors'](self.plan,changed,self.book)))
        finally:changed.close()
        changed=openpyxl.load_workbook(self.source,data_only=False)
        try:
            changed.remove(changed['Formats'])
            self.assertTrue(any(e[0]=='sheet_order' for e in self.mod['preservation_errors'](self.plan,changed,self.book)))
        finally:changed.close()
    def test_stale_plan_does_not_create_snapshot(self):
        p=copy.deepcopy(self.plan);p['input_sha256']='0'*64
        path=self.root/'A1_修改计划.json';path.write_text(json.dumps(p),encoding='utf-8')
        errors,snapshot=self.mod['prepare_source'](1,p,path)
        self.assertEqual(errors[0][0],'source_changed');self.assertIsNone(snapshot)
        self.assertFalse(list(self.root.glob('*快照*')))
    @unittest.skipUnless(os.environ.get('FIGURE_SAVE_INTEGRATION')=='1','requires configured Node/artifact runtime and PowerShell')
    def test_full_shared_pipeline_preserves_native_image_without_copied_scripts(self):
        node=os.environ.get('FIGURE_NODE') or shutil.which('node')
        shell=shutil.which('pwsh') or shutil.which('powershell')
        self.assertTrue(node and shell,'integration explicitly requested but runtime is missing')
        path=self.root/'A1_修改计划.json';path.write_text(json.dumps(self.plan),encoding='utf-8')
        env={**os.environ,'SRDP_TASK_DIR':str(self.root),'FIGURE_PYTHON':sys.executable,'PYTHONIOENCODING':'utf-8'}
        commands=[[node,str(TOOLS/'author.mjs'),'author','1'],
                  [sys.executable,'-B',str(TOOLS/'merge_cells.py'),'1'],
                  [shell,'-NoProfile','-File',str(TOOLS/'recalculate.ps1'),'-PaperIds','1','-BundledEngine'],
                  [node,str(TOOLS/'author.mjs'),'preview','1']]
        for cmd in commands:
            r=subprocess.run(cmd,env=env,capture_output=True,text=True,encoding='utf-8')
            self.assertEqual(r.returncode,0,r.stdout+'\n'+r.stderr)
        saved=openpyxl.load_workbook(self.source,data_only=False)
        self.assertEqual(saved['Data']['A1'].value,3)
        self.assertEqual(saved['Data']['B1'].value,2)
        self.assertEqual(saved['Data']['A1'].number_format,'0.00')
        self.assertEqual(saved.sheetnames,['Data','Formats'])
        self.assertEqual(len(saved['Data']._images),1)
        result=json.loads((self.root/'续补保存检查.json').read_text(encoding='utf-8'))['1']
        self.assertFalse(result['errors'])
        self.assertTrue((self.root/'A1_preview_synthetic.png').is_file())
        self.assertFalse(list(self.root.glob('*.py'))+list(self.root.glob('*.mjs'))+list(self.root.glob('*.ps1')))
        saved.close()

if __name__=='__main__':unittest.main()
