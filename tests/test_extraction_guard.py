"""Retained-reading and agreed-sampling guards; not visual accuracy tests."""
import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import openpyxl

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'tools'))
from audit_workbook import audit


class ExtractionGuardTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.book=self.root/'book.xlsx';self.contract=self.root/'audit.json'
        self.wb=openpyxl.Workbook();self.ws=self.wb.active;self.ws.title='Data'
        self.ws.append([1980,1990,2000]);self.ws.append([-0.41,-1.29,0.21])
        self.rows=[{'sheet':'Data','cell':f'{c}2','raw_value':v,'method':'direct',
                    'source':'synthetic original pixels'} for c,v in zip('ABC',[-.41,-1.29,.21])]

    def run_audit(self,rows=None,mutate_after_hash=False):
        path=self.root/'raw.json'
        path.write_text(json.dumps(self.rows if rows is None else rows),encoding='utf-8')
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        if mutate_after_hash:path.write_text('[]',encoding='utf-8')
        self.wb.save(self.book)
        rules=[{'kind':'extraction_records','sheet':'Data','range':'A2:C2','json':'raw.json',
                'sha256':digest,'source':'retained synthetic readings'},
               {'kind':'sampling','sheet':'Data','range':'A1:C1',
                'expected':[1980,1990,2000],'source':'agreed synthetic sampling'}]
        self.contract.write_text(json.dumps({'rules':rules}),encoding='utf-8')
        return audit(self.book,self.contract)

    def test_real_spike_is_preserved_but_neighbor_mean_fails(self):
        self.assertFalse(self.run_audit()['failures'])
        self.ws['B2']=-.10
        self.assertIn('被替换',str(self.run_audit()['failures']))

    def test_missing_cannot_be_filled_as_direct_or_by_interpolation_over_raw(self):
        rows=copy.deepcopy(self.rows);rows[1]['raw_value']=None
        self.assertTrue(self.run_audit(rows)['failures'])
        rows=copy.deepcopy(self.rows);rows[1].update(method='interpolated',calculation={'value':-.1,'inputs':['A2','C2']})
        self.ws['B2']=-.1
        self.assertIn('覆盖已有',str(self.run_audit(rows)['failures']))

    def test_explicit_missing_and_separate_interpolation(self):
        rows=copy.deepcopy(self.rows);rows[1].update(raw_value=None,method='missing')
        self.ws['B2']=None
        self.assertFalse(self.run_audit(rows)['failures'])
        rows[1].update(method='interpolated',calculation={'value':-.1,'inputs':['A2','C2']})
        self.ws['B2']=-.1
        self.assertFalse(self.run_audit(rows)['failures'])

    def test_reviewed_correction_requires_pixel_and_source(self):
        rows=copy.deepcopy(self.rows);rows[1]['method']='reviewed';self.ws['B2']=-1.3
        self.assertTrue(self.run_audit(rows)['failures'])
        rows[1]['review']={'pixel':[237,924],'source':'synthetic crop','value':-1.3}
        self.assertFalse(self.run_audit(rows)['failures'])

    def test_same_count_wrong_year_and_source_rewrite_fail(self):
        self.ws['B1']=1995
        self.assertIn('取样',str(self.run_audit()['failures']))
        self.ws['B1']=1990
        self.assertIn('摘要',str(self.run_audit(mutate_after_hash=True)['failures']))

    def test_coverage_cannot_omit_changed_cell(self):
        self.ws['B2']=-.1
        self.assertIn('覆盖',str(self.run_audit([self.rows[0],self.rows[2]])['failures']))


if __name__=='__main__':unittest.main()
